import openpyxl
import io
import uuid
import openai
import fitz  # PyMuPDF
from docx import Document
from core.config import settings
from services.storage import storage

# Initialize Kimi client for embeddings
client = openai.OpenAI(
    api_key=settings.KIMI_API_KEY,
    base_url="https://api.moonshot.cn/v1",
)

def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list:
    """Improved chunking logic with overlap."""
    chunks = []
    if not text:
        return chunks
    
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += (chunk_size - overlap)
    return chunks

def process_document(file_content: bytes, filename: str) -> int:
    try:
        ext = filename.split('.')[-1].lower()
        all_text = ""
        qa_entries = []  # For Q&A format Excel rows
        
        if ext == 'xlsx':
            wb = openpyxl.load_workbook(filename=io.BytesIO(file_content), data_only=True)
            for sheet in wb.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue
                
                # Detect Q&A format: check if header row contains question/answer keywords
                header = [str(h).strip().lower() if h else '' for h in rows[0]]
                qa_keywords = {'问题', '答案', 'question', 'answer', 'q', 'a', '问', '答'}
                is_qa = len(set(header) & qa_keywords) >= 2
                
                if is_qa:
                    # Q&A mode: store each row as independent entry
                    for row in rows[1:]:
                        if any(row):
                            vals = [str(v) for v in row if v is not None]
                            if len(vals) >= 2:
                                qa_entry = f"问题: {vals[0]}\n答案: {vals[1]}"
                            else:
                                qa_entry = ", ".join(vals)
                            qa_entries.append(qa_entry)
                else:
                    # Generic mode: concatenate all text for chunking
                    for row in rows:
                        if any(row):
                            all_text += ", ".join([str(val) for val in row if val is not None]) + "\n"
        
        elif ext == 'pdf':
            doc = fitz.open(stream=file_content, filetype="pdf")
            for page in doc:
                all_text += page.get_text() + "\n"
            doc.close()
            
        elif ext in ['doc', 'docx']:
            doc = Document(io.BytesIO(file_content))
            for para in doc.paragraphs:
                all_text += para.text + "\n"
        
        else:
            print(f"Unsupported file format: {ext}")
            return 0

        # Build final chunks: Q&A entries are used as-is, generic text is chunked
        chunks = qa_entries + chunk_text(all_text)
        
        if not chunks:
            return 0

        # Ensure collection exists
        collection_name = "knowledge_base"
        storage.ensure_collection(collection_name, vector_size=1280)

        # Generate embeddings and upsert
        batch_size = 20
        total_inserted = 0
        
        for i in range(0, len(chunks), batch_size):
            batch_chunks = chunks[i:i + batch_size]
            
            res = client.embeddings.create(
                model="embedding-2",
                input=batch_chunks
            )
            
            vectors = [item.embedding for item in res.data]
            ids = [str(uuid.uuid4()) for _ in batch_chunks]
            payloads = [{"text": text, "source": filename} for text in batch_chunks]
            
            success = storage.upsert_documents(collection_name, ids, vectors, payloads)
            if success:
                total_inserted += len(batch_chunks)
        
        return total_inserted
    except Exception as e:
        print(f"Error parsing document {filename}: {e}")
        return 0
